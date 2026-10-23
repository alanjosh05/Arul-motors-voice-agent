"""Load the client's filled data-collection workbook into the CRM database.

    python import_client_data.py Arul-Motors-Customer-Data.xlsx           # check only
    python import_client_data.py Arul-Motors-Customer-Data.xlsx --apply   # write

Sheets used: "Service Types" (the workshop's schedule) and "Customers & Bikes"
(one row per bike). A row with only a mobile number is an extra number for
the customer above it. Re-running is safe: bikes are matched by registration
and updated, service types by name.

The report masks mobile numbers; the workbook itself must never be committed.
"""

import argparse
import re
import sys
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

import database


def _text(value) -> str:
    """Trimmed text with runs of spaces collapsed ('Yamaha  MT 15' -> 'Yamaha MT 15')."""
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _mobile(value) -> str:
    digits = re.sub(r"\D", "", str(value or ""))
    return digits[-10:] if len(digits) >= 10 else digits


def _mask(mobile: str) -> str:
    return f"xxxxxx{mobile[-4:]}" if mobile else "-"


def _day(value) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return database.normalize_date(_text(value)) if value else None


def _whole(value) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def read_service_types(ws) -> list[dict]:
    types = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = _text(row[0])
        if not name:
            continue
        types.append({
            "name": name,
            "description": _text(row[1]) or None,
            "is_free": _text(row[2]).lower().startswith("y"),
            "interval_days": _whole(row[3]),
            "interval_km": _whole(row[4]),
            "price": _whole(row[5]),
            "active": True,
        })
    return types


def read_bikes(ws, problems: list[str]) -> tuple[list[dict], dict[str, list[str]]]:
    """Bike rows, plus extra mobile numbers per customer's main number."""
    bikes, extra_numbers = [], {}
    for number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        name, mobile = _text(row[0]), _mobile(row[1])
        if not any(v not in (None, "") for v in row[:10]):
            continue
        if not name and mobile and not _text(row[3]):
            if bikes:
                extra_numbers.setdefault(bikes[-1]["mobile"], []).append(mobile)
            continue
        bike = {
            "row": number,
            "name": name,
            "mobile": mobile,
            "model": _text(row[2]),
            "registration": _text(row[3]).upper(),
            "purchase_date": _day(row[4]),
            "last_service_date": _day(row[5]),
            "last_service_type": _text(row[6]),
            "odometer_km": _whole(row[7]),
            "next_service_date": _day(row[8]),
            "next_service_type": _text(row[9]),
        }
        missing = [k for k in ("name", "model", "registration", "purchase_date") if not bike[k]]
        if len(mobile) != 10:
            missing.append("10-digit mobile")
        if missing:
            problems.append(f"Row {number}: skipped, missing {', '.join(missing)}.")
            continue
        bikes.append(bike)
    return bikes, extra_numbers


def plan_next_service(bike: dict, types: list[dict], problems: list[str]) -> bool:
    """Fill or sanity-check the next service. Returns False to skip the bike."""
    by_name = {t["name"].lower(): t for t in types}
    where = f"Row {bike['row']} ({bike['registration']})"
    for key in ("last_service_type", "next_service_type"):
        if bike[key] and bike[key].lower() not in by_name:
            problems.append(f"{where}: unknown service type '{bike[key]}', skipped.")
            return False
        if bike[key]:
            bike[key] = by_name[bike[key].lower()]["name"]

    if not bike["next_service_date"] or not bike["next_service_type"]:
        # Work it out from the schedule: the type after the last one, due its
        # interval after the last service (or the purchase).
        scheduled = [t for t in types if t["interval_days"]]
        names = [t["name"] for t in scheduled]
        if bike["last_service_type"] in names:
            following = scheduled[min(names.index(bike["last_service_type"]) + 1, len(scheduled) - 1)]
        else:
            following = scheduled[0]
        start = date.fromisoformat(bike["last_service_date"] or bike["purchase_date"])
        bike["next_service_type"] = bike["next_service_type"] or following["name"]
        due = start + timedelta(days=following["interval_days"])
        bike["next_service_date"] = bike["next_service_date"] or due.isoformat()
        problems.append(f"{where}: next service worked out as {bike['next_service_type']} on {bike['next_service_date']}.")

    last = bike["last_service_date"] or bike["purchase_date"]
    if bike["next_service_date"] <= last:
        # e.g. 17-01-2026 typed for 17-01-2027: move it to the next year after the last service.
        fixed = date.fromisoformat(bike["next_service_date"])
        while fixed.isoformat() <= last:
            fixed = fixed.replace(year=fixed.year + 1)
        problems.append(
            f"{where}: next service {bike['next_service_date']} is before the last service "
            f"({last}); corrected to {fixed.isoformat()}. Please confirm with the showroom."
        )
        bike["next_service_date"] = fixed.isoformat()
    return True


def dedupe(bikes: list[dict], problems: list[str]) -> list[dict]:
    """Same registration twice: keep the row with the latest last service."""
    keep: dict[str, dict] = {}
    for bike in bikes:
        other = keep.get(bike["registration"])
        if not other:
            keep[bike["registration"]] = bike
            continue
        newer = max(other, bike, key=lambda b: b["last_service_date"] or "")
        problems.append(
            f"Rows {other['row']} and {bike['row']}: {bike['registration']} listed twice "
            f"(last service {other['last_service_date']} / {bike['last_service_date']}); "
            f"kept row {newer['row']}. Please confirm which is right."
        )
        keep[bike["registration"]] = newer
    return list(keep.values())


def apply_service_types(types: list[dict]) -> None:
    existing = {t["name"].lower(): t for t in database.list_service_types()}
    for t in types:
        current = existing.pop(t["name"].lower(), None)
        if current:
            ok, message = database.update_service_type(current["id"], **t)
        else:
            ok, message, _ = database.create_service_type(**t)
        print(f"  {'ok ' if ok else 'ERR'} service type {t['name']}: {message}")
    # Defaults the workshop doesn't use: delete if unused, otherwise hide.
    for t in existing.values():
        ok, message = database.delete_service_type(t["id"])
        if not ok:
            database.update_service_type(t["id"], active=False)
            message = "in use, deactivated"
        print(f"  --  removed default type {t['name']}: {message}")
    # Schedule order as on the sheet.
    order = {t["name"].lower(): i for i, t in enumerate(types)}
    ids = sorted(database.list_service_types(), key=lambda t: order.get(t["name"].lower(), 999))
    conn = database._connect()
    for position, t in enumerate(ids, start=1):
        conn.execute("UPDATE service_types SET sort_order = ? WHERE id = ?", (position, t["id"]))
    conn.commit()
    conn.close()


def apply_bikes(bikes: list[dict]) -> None:
    vehicles = {v["registration_number"]: v for v in database.list_vehicles()}
    for b in bikes:
        fields = {
            "model": b["model"],
            "purchase_date": b["purchase_date"],
            "last_service_date": b["last_service_date"],
            "odometer_km": b["odometer_km"],
            "next_service_date": b["next_service_date"],
            "next_service_type": b["next_service_type"],
        }
        if b["registration"] in vehicles:
            ok, message = database.update_vehicle(vehicles[b["registration"]]["id"], **fields)
        else:
            ok, message = database.add_customer_vehicle(
                name=b["name"], mobile_number=b["mobile"], registration_number=b["registration"], **fields
            )
        print(f"  {'ok ' if ok else 'ERR'} {b['registration']} {b['name']}: {message}")
    record_last_services(bikes)


def record_last_services(bikes: list[dict]) -> None:
    """Service history: each bike's last service as a COMPLETED booking, with
    the type from the sheet (init_db would otherwise guess it from the next one)."""
    conn = database._connect()
    cursor = conn.cursor()
    added = 0
    for b in bikes:
        if not b["last_service_date"]:
            continue
        cursor.execute("SELECT id, customer_id FROM vehicles WHERE registration_number = ?", (b["registration"],))
        vehicle = cursor.fetchone()
        cursor.execute(
            "SELECT 1 FROM service_appointments WHERE vehicle_id = ? AND status = 'COMPLETED' AND appointment_date = ?",
            (vehicle["id"], b["last_service_date"]),
        )
        if cursor.fetchone():
            continue
        cursor.execute(
            "INSERT INTO service_appointments (id, vehicle_id, customer_id, appointment_date, service_type, notes, status) "
            "VALUES (?, ?, ?, ?, ?, ?, 'COMPLETED')",
            (
                database._next_appointment_id(cursor),
                vehicle["id"],
                vehicle["customer_id"],
                b["last_service_date"],
                b["last_service_type"] or "Service",
                "Imported from the showroom's records.",
            ),
        )
        added += 1
    conn.commit()
    conn.close()
    print(f"  ok  service history: {added} past service(s) recorded")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("workbook")
    parser.add_argument("--apply", action="store_true", help="write to the database (default: check only)")
    args = parser.parse_args()

    wb = load_workbook(args.workbook, data_only=True)
    problems: list[str] = []
    types = read_service_types(wb["Service Types"])
    if not any(t["interval_days"] for t in types):
        print("The Service Types sheet has no scheduled types; nothing imported.")
        return 1
    bikes, extra_numbers = read_bikes(wb["Customers & Bikes"], problems)
    bikes = [b for b in dedupe(bikes, problems) if plan_next_service(b, types, problems)]

    customers: dict[str, str] = {}
    for b in bikes:
        if customers.setdefault(b["mobile"], b["name"]) != b["name"]:
            problems.append(
                f"Row {b['row']}: mobile {_mask(b['mobile'])} is also used by {customers[b['mobile']]}; "
                f"the bike goes to that customer."
            )

    print(f"Service types: {len(types)}")
    for t in types:
        interval = f"every {t['interval_days']} days / {t['interval_km']} km" if t["interval_days"] else "on request"
        print(f"  - {t['name']} ({'free' if t['is_free'] else 'paid'}, {interval})")
    print(f"Customers: {len(customers)}   Bikes: {len(bikes)}")
    if extra_numbers:
        count = sum(len(v) for v in extra_numbers.values())
        print(f"Extra mobile numbers: {count} for {len(extra_numbers)} customers "
              "(not stored: the CRM keeps one number per customer; the first one is called).")
    print(f"\nNotes ({len(problems)}):" if problems else "\nNo problems found.")
    for p in problems:
        print(f"  - {p}")

    if not args.apply:
        print("\nCheck only. Run again with --apply to write to", database.DB_FILE)
        return 0
    database.init_db()
    print("\nWriting to", database.DB_FILE)
    apply_service_types(types)
    apply_bikes(bikes)
    print("Done.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
