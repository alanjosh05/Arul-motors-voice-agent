"""Excel export of the CRM (Settings > Data & backups > Export to Excel).

Five sheets the showroom can filter and work on: Customers & Bikes (one row
per bike), Service Records (what was done at each service), Bookings, Calls
and Service Types. Service report PDFs stay in the CRM. Dates are real Excel dates and
mobile numbers stay text.
"""

import io
from datetime import date, datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import database

OUTCOME_LABELS = {
    "date_confirmed": "Date confirmed",
    "callback_requested": "Call back",
    "not_interested": "Not interested",
    "already_serviced": "Already serviced",
    "bike_sold": "Bike sold",
    "wrong_number": "Wrong number",
    "no_answer": "No answer",
}
STATE_LABELS = {"upcoming": "Upcoming", "completed": "Serviced", "cancelled": "Cancelled", "missed": "Missed"}

HEADER = PatternFill("solid", fgColor="1F3A8A")
DAY = "DD-MM-YYYY"
DAY_TIME = "DD-MM-YYYY HH:MM"


def _day(value) -> date | None:
    try:
        return date.fromisoformat(str(value)[:10]) if value else None
    except ValueError:
        return None


def _local_time(value) -> datetime | None:
    """Database timestamps are UTC ('YYYY-MM-DD HH:MM:SS'): show IST, no tzinfo
    (Excel can't store time zones)."""
    if not value:
        return None
    try:
        moment = value if isinstance(value, datetime) else datetime.fromisoformat(str(value))
    except ValueError:
        return None
    if moment.tzinfo is None:
        moment = moment.replace(tzinfo=timezone.utc)
    return moment.astimezone(database.SHOWROOM_TZ).replace(tzinfo=None)


def _sheet(wb: Workbook, title: str, columns: list[tuple[str, int, str | None]], rows: list[list]) -> None:
    """columns: (header, width, number format or None)."""
    ws = wb.create_sheet(title)
    ws.append([c[0] for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    for i, (_, width, fmt) in enumerate(columns, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = width
        if fmt:
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                cell.number_format = fmt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(len(rows) + 1, 2)}"


def build_workbook() -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    customers = database.list_customers()
    bike_rows = []
    for c in sorted(customers, key=lambda c: c["name"].lower()):
        booking = c.get("next_appointment")
        call = c.get("last_call")
        shared = [
            c["name"],
            c["mobile_number"],
            c.get("alternate_mobile") or "",
        ]
        tail = [
            _day(booking["appointment_date"]) if booking else None,
            _local_time(call["created_at"]) if call else None,
            OUTCOME_LABELS.get(call["outcome"], call["outcome"]) if call else "",
        ]
        for v in c["vehicles"] or [None]:
            bike = (
                [v["model"], v["registration_number"],
                 (v.get("dealer_name") or "Another dealer") if v.get("bought_elsewhere") else "Arul Motors",
                 _day(v["purchase_date"]), _day(v["last_service_date"]),
                 v["odometer_km"], _day(v["next_service_date"]), v["next_service_type"]]
                if v else ["", "", "", None, None, None, None, ""]
            )
            bike_rows.append(shared + bike + tail)
    _sheet(wb, "Customers & Bikes", [
        ("Customer", 24, None), ("Mobile", 13, "@"), ("Alternate mobile", 15, "@"),
        ("Bike model", 30, None), ("Registration", 15, None), ("Bought from", 18, None), ("Purchase date", 13, DAY),
        ("Last service", 13, DAY), ("Odometer (km)", 13, "#,##0"), ("Next service due", 15, DAY),
        ("Next service type", 22, None), ("Next booking", 13, DAY), ("Last call", 16, DAY_TIME),
        ("Last call outcome", 17, None),
    ], bike_rows)

    def yes_no(value):
        return "" if value is None else ("Yes" if value else "No")

    _sheet(wb, "Service Records", [
        ("Delivered", 12, DAY), ("Job card", 10, None), ("Customer", 22, None), ("Mobile", 13, "@"),
        ("Bike model", 28, None), ("Registration", 15, None), ("Frame no.", 20, None), ("Engine no.", 16, None),
        ("Type of job", 15, None), ("Service", 22, None), ("Received", 12, DAY), ("Odometer (km)", 13, "#,##0"),
        ("Customer requests", 35, None), ("Jobs done", 45, None), ("Other work", 35, None), ("Parts", 55, None),
        ("Parts ₹", 10, "#,##0"), ("Labour", 35, None), ("Labour ₹", 10, "#,##0"), ("Total ₹", 10, "#,##0"),
        ("Delivery checks", 45, None), ("Feedback", 12, None), ("On time", 8, None), ("Washing OK", 10, None),
        ("Staff OK", 9, None), ("Charges fair", 11, None), ("Repair OK", 10, None),
        ("Technician", 15, None), ("Mechanic", 15, None), ("Supervisor", 15, None), ("Remarks", 35, None),
        ("Next service", 12, DAY), ("Next service type", 22, None), ("Job card PDF", 24, None), ("Closed by", 15, None),
    ], [
        [_day(r["service_date"]), r["job_card_no"] or "", r["customer_name"], r["mobile_number"], r["model"],
         r["registration_number"], r["frame_number"] or "", r["engine_number"] or "", r["job_type"] or "",
         r["service_type"], _day(r["received_date"]), r["odometer_km"], r["customer_requests"] or "",
         ", ".join(r["jobs_done"]), r["work_done"] or "", database.parts_summary(r["parts"]), r["parts_total"],
         "; ".join(f"{x['description']} ₹{x['amount']}" for x in r["labour"]), r["labour_total"], r["amount"],
         ", ".join(r["delivery_checks"]), r["feedback"].get("rating", ""),
         yes_no(r["feedback"].get("delivery_on_time")), yes_no(r["feedback"].get("washing_quality")),
         yes_no(r["feedback"].get("staff_behaviour")), yes_no(r["feedback"].get("fairness_of_charges")),
         yes_no(r["feedback"].get("repairing_quality")),
         r["technician"] or "", r["mechanic"] or "", r["supervisor"] or "", r["remarks"] or "",
         _day(r["next_service_date"]), r["next_service_type"] or "", r["document_name"] or "", r["created_by"] or ""]
        for r in database.list_service_records()
    ])

    bookings = database.search_appointments()
    _sheet(wb, "Bookings", [
        ("Booking", 10, None), ("Date", 12, DAY), ("Status", 11, None), ("Customer", 24, None),
        ("Mobile", 13, "@"), ("Bike model", 30, None), ("Registration", 15, None),
        ("Service type", 22, None), ("Notes", 40, None), ("Created", 16, DAY_TIME),
    ], [
        [b["id"], _day(b["appointment_date"]), STATE_LABELS.get(b["state"], b["state"]), b["customer_name"],
         b["mobile_number"], b["model"], b["registration_number"], b["service_type"], b["notes"] or "",
         _local_time(b["created_at"])]
        for b in bookings
    ])

    calls = database.get_call_logs()
    _sheet(wb, "Calls", [
        ("When", 16, DAY_TIME), ("Customer", 24, None), ("Mobile", 13, "@"), ("Bike", 30, None),
        ("Registration", 15, None), ("Outcome", 17, None), ("Call back on", 13, DAY), ("Notes", 50, None),
    ], [
        [_local_time(c["created_at"]), c["customer_name"], c["mobile_number"], c["model"] or "",
         c["registration_number"] or "", OUTCOME_LABELS.get(c["outcome"], c["outcome"]),
         _day(c["callback_date"]), c["notes"] or ""]
        for c in calls
    ])

    _sheet(wb, "Service Types", [
        ("Service", 26, None), ("Free?", 7, None), ("Due after (days)", 15, None), ("Due after (km)", 14, "#,##0"),
        ("Price (₹)", 10, "#,##0"), ("Overdue after (days)", 18, None), ("Active", 8, None), ("Description", 50, None),
    ], [
        [t["name"], "Yes" if t["is_free"] else "No", t["interval_days"], t["interval_km"], t["price"],
         t["overdue_days"], "Yes" if t["active"] else "No", t["description"] or ""]
        for t in database.list_service_types()
    ])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
